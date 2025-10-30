#!/usr/bin/env python3
"""
Conversor de JSON a Word - Reconstruye documentos desde JSON extraído
"""

import json
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import re

def add_heading_with_format(doc, text, level=1):
    """Agrega un título con formato personalizado"""
    heading = doc.add_heading(level=level)
    run = heading.add_run(text)
    run.font.size = Pt(14 if level == 1 else 12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return heading

def add_formatted_paragraph(doc, text, bold=False, italic=False, size=11):
    """Agrega un párrafo con formato"""
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    return paragraph

def process_content(doc, content):
    """Procesa el contenido y lo formatea apropiadamente"""
    # Dividir el contenido en líneas
    lines = content.split('\n')

    for line in lines:
        line = line.strip()

        if not line:
            continue

        # Detectar y formatear títulos de capítulos (ej: "01. CAPITULO 100 PRELIMINARES")
        if re.match(r'^\d+\.\s*CAPITULO\s+\d+', line, re.IGNORECASE):
            add_heading_with_format(doc, line, level=1)

        # Detectar secciones (ej: "01.01. SEC. 103 MANTENIMIENTO...")
        elif re.match(r'^\d+\.\d+\.\s*SEC', line, re.IGNORECASE) or re.match(r'^SECCIÓN\s+\d+', line, re.IGNORECASE):
            add_heading_with_format(doc, line, level=2)

        # Detectar subsecciones numeradas (ej: "103.1 Descripción")
        elif re.match(r'^\d+\.\d+\s+[A-Za-z]', line):
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(line)
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102)

        # Detectar listas con guiones
        elif line.startswith('-'):
            paragraph = doc.add_paragraph(line[1:].strip(), style='List Bullet')
            paragraph.paragraph_format.left_indent = Inches(0.5)

        # Detectar listas numeradas
        elif re.match(r'^\d+[\.)]\s+', line):
            # Extraer el texto después del número
            text = re.sub(r'^\d+[\.)]\s+', '', line)
            paragraph = doc.add_paragraph(text, style='List Number')
            paragraph.paragraph_format.left_indent = Inches(0.5)

        # Detectar letras de listas (a., b., c.)
        elif re.match(r'^[a-z][\.)]\s+', line, re.IGNORECASE):
            text = re.sub(r'^[a-z][\.)]\s+', '', line, flags=re.IGNORECASE)
            paragraph = doc.add_paragraph(text, style='List Bullet')
            paragraph.paragraph_format.left_indent = Inches(0.75)

        # Párrafos normales
        else:
            paragraph = doc.add_paragraph(line)
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_after = Pt(6)

def create_table_from_text(doc, table_text):
    """Crea una tabla en el documento si detecta formato de tabla"""
    lines = table_text.strip().split('\n')
    if len(lines) < 2:
        return False

    # Detectar si hay un separador de tabla
    separator_line = None
    for i, line in enumerate(lines):
        if '|' in line and '-' in line:
            separator_line = i
            break

    if separator_line is None:
        return False

    # Procesar las filas
    header_line = lines[separator_line - 1] if separator_line > 0 else None
    data_lines = lines[separator_line + 1:]

    if not header_line:
        return False

    # Contar columnas
    header_cells = [cell.strip() for cell in header_line.split('|') if cell.strip()]
    num_cols = len(header_cells)
    num_rows = len([line for line in data_lines if '|' in line]) + 1  # +1 para header

    if num_cols == 0 or num_rows == 0:
        return False

    # Crear tabla
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = 'Light Grid Accent 1'

    # Agregar encabezado
    for i, cell_text in enumerate(header_cells):
        cell = table.rows[0].cells[i]
        cell.text = cell_text
        cell.paragraphs[0].runs[0].font.bold = True

    # Agregar datos
    row_idx = 1
    for line in data_lines:
        if '|' not in line:
            continue
        cells = [cell.strip() for cell in line.split('|') if cell.strip()]
        if len(cells) == num_cols:
            for i, cell_text in enumerate(cells):
                table.rows[row_idx].cells[i].text = cell_text
            row_idx += 1

    return True

def json_to_word(json_path, output_path):
    """Convierte un archivo JSON a documento Word"""

    print(f"📄 Leyendo archivo JSON: {json_path}")
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    # Crear documento
    doc = Document()

    # Configurar estilos predeterminados
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Agregar título principal
    title = doc.add_heading(level=0)
    run = title.add_run(data['metadata']['subject'])
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Agregar información de metadatos
    doc.add_paragraph()
    metadata_para = doc.add_paragraph()
    metadata_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = metadata_para.add_run(f"Total de páginas: {data['metadata']['total_pages']}")
    run.font.italic = True
    run.font.size = Pt(10)

    doc.add_page_break()

    # Agregar contenido página por página
    print(f"📝 Procesando {len(data['pages'])} páginas...")

    current_content = ""
    in_table = False
    table_lines = []

    for page in data['pages']:
        if not page.get('has_content', False):
            continue

        content = page['content']

        # Verificar si hay tablas en el contenido
        if '|' in content and '-' in content:
            # Procesar línea por línea para detectar tablas
            lines = content.split('\n')
            i = 0
            while i < len(lines):
                line = lines[i]

                # Detectar inicio de tabla
                if '|' in line and not in_table:
                    # Ver si la siguiente línea es un separador
                    if i + 1 < len(lines) and '|' in lines[i + 1] and '-' in lines[i + 1]:
                        in_table = True
                        table_lines = [line]
                        i += 1
                        continue

                # Si estamos en una tabla
                if in_table:
                    table_lines.append(line)

                    # Verificar si la tabla terminó
                    if '|' not in line or (i + 1 < len(lines) and '|' not in lines[i + 1]):
                        # Crear la tabla
                        table_text = '\n'.join(table_lines)
                        if not create_table_from_text(doc, table_text):
                            # Si no se pudo crear como tabla, agregar como texto normal
                            process_content(doc, table_text)

                        doc.add_paragraph()
                        in_table = False
                        table_lines = []
                    i += 1
                    continue

                # Procesar contenido normal
                process_content(doc, line)
                i += 1
        else:
            # No hay tablas, procesar normalmente
            process_content(doc, content)

        # Agregar un salto de párrafo entre páginas
        if page['page_number'] < data['metadata']['total_pages']:
            doc.add_paragraph()

    # Guardar documento
    print(f"💾 Guardando documento Word: {output_path}")
    doc.save(output_path)
    print(f"✅ Documento generado exitosamente!")

    return output_path

def main():
    """Función principal"""
    # Rutas
    json_path = "output/content/eett_vias_content.json"
    output_dir = "output/word"

    # Crear directorio de salida
    os.makedirs(output_dir, exist_ok=True)

    # Generar nombre de archivo de salida
    output_filename = "EETT_VIAS.docx"
    output_path = os.path.join(output_dir, output_filename)

    # Verificar que existe el archivo JSON
    if not os.path.exists(json_path):
        print(f"❌ Error: No se encuentra el archivo JSON: {json_path}")
        return

    print("=" * 60)
    print("📚 CONVERSOR JSON A WORD")
    print("=" * 60)

    # Convertir
    json_to_word(json_path, output_path)

    print("\n" + "=" * 60)
    print(f"✅ Archivo generado: {output_path}")
    print("=" * 60)

if __name__ == "__main__":
    main()
